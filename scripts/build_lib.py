"""Helpers to write formatted Excel workbooks and Markdown documents."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import OUT_DIR

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NOTE_FONT = Font(italic=True, color="7A2E00")

SIM_NOTE = ("All counts and values in simulated sheets describe FICTIONAL data created for analytical testing. "
            "They are not VELOOP users, earnings, withdrawals, revenue, retention, conversion or performance.")


def write_xlsx(name, title, purpose, sheets, simulated=False, notes=()):
    """sheets: list of (sheet_name, DataFrame). Adds a README sheet first."""
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    rows = [("VELOOP Rewards: Data Analyst Task 01", ""), ("Deliverable", title), ("Purpose", purpose),
            ("Status", "Proposed analytics specification. Nothing here is verified as implemented in VELOOP production."),
            ("Prepared", "2026-10-01"), ("Sheets", ", ".join(s for s, _ in sheets))]
    if simulated:
        rows.append(("SIMULATED DATA", SIM_NOTE))
    for n in notes:
        rows.append(("Note", n))
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for r in range(2, len(rows) + 1):
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if ws.cell(r, 1).value == "SIMULATED DATA":
            ws.cell(r, 1).font = Font(bold=True, color="B00020")
            ws.cell(r, 2).font = NOTE_FONT
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 120
    for sheet_name, df in sheets:
        s = wb.create_sheet(sheet_name[:31])
        df = df.copy()
        for c in df.columns:
            df[c] = df[c].map(lambda v: "; ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v)
        s.append(list(df.columns))
        for row in df.itertuples(index=False):
            s.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
        for i, col in enumerate(df.columns, 1):
            cell = s.cell(1, i)
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
            s.column_dimensions[get_column_letter(i)].width = max(10, min(60, longest + 2))
        for row in s.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "B2"
        s.auto_filter.ref = s.dimensions
    path = OUT_DIR / name
    wb.save(path)
    return path


def md_table(df, cols=None, max_rows=None):
    df = df if cols is None else df[cols]
    if max_rows:
        df = df.head(max_rows)

    def cell(v):
        if isinstance(v, (list, tuple, set)):
            v = ", ".join(map(str, v))
        if v is None or (isinstance(v, float) and pd.isna(v)):
            v = ""
        return str(v).replace("|", "\\|").replace("\n", " ")
    head = "| " + " | ".join(map(str, df.columns)) + " |"
    sep = "| " + " | ".join("---" for _ in df.columns) + " |"
    body = ["| " + " | ".join(cell(v) for v in row) + " |" for row in df.itertuples(index=False)]
    return "\n".join([head, sep] + body)


def write_md(name, text):
    path = OUT_DIR / name
    path.write_text(text.strip() + "\n")
    return path
